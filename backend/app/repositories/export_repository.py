from __future__ import annotations

import time
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

from app.core.db_engine import DuckDBEngine
from app.core.fc_import_norm_insert import sanitize_text_expr as _sanitize_text_expr
from app.core.fc_import_schema import FC_SCHEMAS, norm_table_name
from app.core.import_count_semantics import (
    CLEANING_COUNTS_VERSION,
    IMPORT_COUNTS_VERSION,
    is_cleaning_success_status,
    is_import_success_status,
    known_public_import_count,
)
from app.core.storage import CaseStorage
from app.domain.controlled_artifact_gate import require_controlled_artifact_publication
from app.repositories.export_csv_copy import build_csv_copy_sql, split_alias
from app.utils.fs import safe_fs_name


EXPORT_TABLES: Sequence[Tuple[str, str]] = (
    ("fc_coercive_measure", "强制措施信息"),
    ("fc_person_contact", "人员联系方式信息"),
    ("fc_person_address", "人员住址信息"),
    ("fc_person", "人员信息"),
    ("fc_sub_account", "关联子账户信息"),
    ("fc_account", "账户信息"),
    ("fc_transaction", "交易明细信息"),
    ("fc_task_fail", "任务信息(失败)"),
    ("fc_task_success", "任务信息(成功)"),
)
EXPORT_FOLDER_SUFFIX = "（数据信息）"
ROW_LIMIT_XLSX = 1_048_576
CHUNK_SIZE = 2000

class ExportCancelled(Exception):
    pass

class ExportRepository:
    """Backend export executor aligned with the persisted export workflow."""

    def __init__(self, storage: Optional[CaseStorage] = None) -> None:
        self._storage = storage or CaseStorage()

    @property
    def storage(self) -> CaseStorage:
        return self._storage

    def case_exists(self, case_id: str) -> bool:
        return self._storage.get_case(case_id) is not None

    def ensure_cleaned_export_ready(self, case_id: str) -> None:
        self._storage.ensure_funds_tables(case_id)
        con = DuckDBEngine(self._storage.case_db(case_id))
        try:
            if not self._table_exists(con, "import_file_log"):
                raise ValueError("cleaned dataset not ready: import count provenance unavailable")

            required_columns = {
                "kind",
                "rows_imported_norm",
                "import_counts_version",
                "status",
                "finished_at",
                "cleaned_status",
                "cleaned_finished_at",
                "cleaned_rows_affected",
                "cleaning_counts_version",
            }
            available_columns = {
                str(row[0])
                for row in con.query(
                    "SELECT column_name FROM information_schema.columns "
                    "WHERE table_schema='main' AND table_name='import_file_log'"
                )
            }
            if not required_columns.issubset(available_columns):
                raise ValueError("cleaned dataset not ready: import count provenance unavailable")

            rows = con.query(
                "SELECT kind, rows_imported_norm, import_counts_version, status, finished_at, "
                "cleaned_status, cleaned_finished_at, cleaned_rows_affected, cleaning_counts_version "
                "FROM import_file_log "
                "WHERE case_id=? AND kind IN ('fc_transaction', 'fc_account')",
                (case_id,),
            )
            scoped_rows: list[tuple] = []
            for row in rows:
                kind, rows_imported_norm, import_version, status, finished_at = row[:5]
                if not is_import_success_status(status) or not self._has_completed_at(finished_at):
                    raise ValueError("cleaned dataset not ready: import scope file(s) not completed")
                if type(import_version) is not int or import_version != IMPORT_COUNTS_VERSION:
                    raise ValueError("cleaned dataset not ready: import count provenance unavailable")
                normalized_count = known_public_import_count(rows_imported_norm)
                if normalized_count is None:
                    raise ValueError("cleaned dataset not ready: import count provenance unavailable")
                if normalized_count > 0:
                    scoped_rows.append(row)

            if not any(str(row[0]) == "fc_transaction" for row in scoped_rows):
                raise ValueError("cleaned dataset not ready: no transaction data in cleaning scope")

            pending_count = 0
            for row in scoped_rows:
                cleaned_status, cleaned_finished_at, cleaned_rows_affected, cleaning_version = row[5:]
                if not is_cleaning_success_status(cleaned_status) or not self._has_completed_at(
                    cleaned_finished_at
                ):
                    pending_count += 1
                    continue
                if (
                    type(cleaning_version) is not int
                    or cleaning_version != CLEANING_COUNTS_VERSION
                    or known_public_import_count(cleaned_rows_affected) is None
                ):
                    raise ValueError("cleaned dataset not ready: cleaning count provenance unavailable")

            if pending_count > 0:
                raise ValueError(
                    f"cleaned dataset not ready: {pending_count} cleaning scope file(s) not completed"
                )
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("cleaned dataset not ready: import count provenance unavailable") from exc
        finally:
            try:
                con.close()
            except Exception:
                pass

    @staticmethod
    def _has_completed_at(value: object) -> bool:
        return type(value) is str and bool(value.strip())

    def run_export(
        self,
        *,
        case_id: str,
        mode: str,
        export_format: str,
        output_name: str = "",
        target_dir: str = "",
        filters: Optional[dict] = None,
        cancel_check: Optional[Callable[[], bool]] = None,
        progress_cb: Optional[Callable[[int, str, Optional[str], Optional[dict]], None]] = None,
    ) -> Dict[str, Any]:
        require_controlled_artifact_publication()
        mode_value = (mode or "").strip().lower()
        if mode_value not in {"raw", "cleaned"}:
            raise ValueError("mode must be raw or cleaned")
        format_value = (export_format or "").strip().lower()
        if format_value not in {"csv", "xlsx"}:
            raise ValueError("export_format must be csv or xlsx")

        case = self._storage.get_case(case_id)

        def _check_cancel() -> None:
            if cancel_check and cancel_check():
                raise ExportCancelled()

        def _emit_progress(
            value: int,
            message: str,
            *,
            stage: Optional[str] = None,
            counters: Optional[dict] = None,
        ) -> None:
            if progress_cb:
                progress_cb(max(0, min(100, int(value))), message, stage, counters or {})

        _check_cancel()

        use_cleaned = mode_value == "cleaned"
        base_name = (output_name or "").strip() or ((case.name if case else "") or case_id)
        safe_name = safe_fs_name(base_name, "未命名案件")
        stamp = datetime.now().strftime("%Y-%m-%d")
        export_tag = "（已清洗）" if use_cleaned else "（原始）"
        folder_name = safe_fs_name(
            f"{stamp} - {safe_name}{EXPORT_FOLDER_SUFFIX}{export_tag}",
            f"{stamp}-export",
        )
        base_dir_text = (target_dir or "").strip()
        export_base = Path(base_dir_text).expanduser() if base_dir_text else self._storage.case_dir(case_id) / "exports"
        if export_base.exists() and not export_base.is_dir():
            raise ValueError("target_dir must be a directory path")
        export_base.mkdir(parents=True, exist_ok=True)
        export_dir = export_base / folder_name
        export_dir.mkdir(parents=True, exist_ok=True)

        label_to_table = {label: table for table, label in EXPORT_TABLES}
        table_to_label = {table: label for table, label in EXPORT_TABLES}
        requested_tables = self._resolve_requested_tables(filters, label_to_table, table_to_label)
        table_specs = [(table, table_to_label[table]) for table in requested_tables]
        if not table_specs:
            raise ValueError("no export tables selected")

        start_ts = time.perf_counter()
        csv_exports: List[Path] = []
        empty_tables: List[str] = []
        created_files: List[Path] = []
        table_stats: Dict[str, dict] = {}

        con = None
        try:
            con = DuckDBEngine(self._storage.case_db(case_id))
            totals: Dict[str, int] = {}
            done: Dict[str, int] = {}
            complete: Dict[str, bool] = {}
            meta: Dict[str, dict] = {}

            for table, label in table_specs:
                _check_cancel()
                table_name = self._resolve_table_name(con, table)
                exists = self._table_exists(con, table_name)
                has_case = self._has_case_id(con, table_name) if exists else False
                total_rows = self._count_rows(con, table_name, has_case, case_id=case_id) if exists else 0
                cols, headers, db_cols, type_map = self._table_columns(
                    con=con,
                    table=table_name,
                    base_table=table,
                    use_cleaned=use_cleaned,
                    case_id=case_id,
                )
                meta[table] = {
                    "label": label,
                    "exists": exists,
                    "has_case": has_case,
                    "total": total_rows,
                    "cols": cols,
                    "headers": headers,
                    "db_cols": db_cols,
                    "type_map": type_map,
                    "table_name": table_name,
                    "base_table": table,
                }
                totals[table] = total_rows
                done[table] = 0
                complete[table] = False
                table_stats[table] = {
                    "label": label,
                    "rows_total": total_rows,
                    "status": "queued",
                }

            self._update_overall_progress(
                done=done,
                totals=totals,
                complete=complete,
                table_count=len(table_specs),
                emit=_emit_progress,
                message="准备导出",
            )

            for table, label in table_specs:
                _check_cancel()
                info = meta[table]
                total_rows = int(info["total"] or 0)
                exists = bool(info["exists"])
                has_case = bool(info["has_case"])
                cols = info["cols"]
                headers = info["headers"] or cols
                type_map = info["type_map"]
                table_name = info["table_name"]
                base_table = info["base_table"]
                safe_label = safe_fs_name(label, "数据表")

                if not exists or total_rows <= 0:
                    empty_tables.append(label)
                    done[table] = 0
                    complete[table] = True
                    table_stats[table] = {
                        "label": label,
                        "rows_total": total_rows,
                        "rows_exported": 0,
                        "status": "empty",
                    }
                    self._update_overall_progress(
                        done=done,
                        totals=totals,
                        complete=complete,
                        table_count=len(table_specs),
                        emit=_emit_progress,
                        message=f"{label}：空表（未导出）",
                        stage="table_empty",
                        counters={"table": table, "rows_total": total_rows},
                    )
                    continue

                use_csv = format_value == "csv" or (base_table == "fc_transaction" and total_rows > ROW_LIMIT_XLSX)
                file_path = export_dir / f"{safe_label}.{'csv' if use_csv else 'xlsx'}"
                created_files.append(file_path)
                table_stats[table]["status"] = "running"
                table_stats[table]["file"] = str(file_path)
                table_stats[table]["format"] = "csv" if use_csv else "xlsx"

                self._update_overall_progress(
                    done=done,
                    totals=totals,
                    complete=complete,
                    table_count=len(table_specs),
                    emit=_emit_progress,
                    message=f"导出中：{label}",
                    stage="table_start",
                    counters={
                        "table": table,
                        "rows_total": total_rows,
                        "format": "csv" if use_csv else "xlsx",
                    },
                )

                if use_csv:
                    rows_written = self._export_csv(
                        con=con,
                        case_id=case_id,
                        table=table_name,
                        has_case=has_case,
                        cols=cols,
                        headers=headers,
                        db_cols=info["db_cols"],
                        type_map=type_map,
                        base_table=base_table,
                        use_cleaned=use_cleaned,
                        file_path=file_path,
                        on_chunk=lambda written: self._on_chunk(
                            written=written,
                            table=table,
                            table_label=label,
                            done=done,
                            totals=totals,
                            complete=complete,
                            table_count=len(table_specs),
                            emit=_emit_progress,
                            check_cancel=_check_cancel,
                        ),
                    )
                    csv_exports.append(file_path)
                else:
                    rows_written = self._export_xlsx(
                        con=con,
                        case_id=case_id,
                        table=table_name,
                        has_case=has_case,
                        cols=cols,
                        headers=headers,
                        db_cols=info["db_cols"],
                        type_map=type_map,
                        base_table=base_table,
                        use_cleaned=use_cleaned,
                        file_path=file_path,
                        on_chunk=lambda written: self._on_chunk(
                            written=written,
                            table=table,
                            table_label=label,
                            done=done,
                            totals=totals,
                            complete=complete,
                            table_count=len(table_specs),
                            emit=_emit_progress,
                            check_cancel=_check_cancel,
                        ),
                    )

                done[table] = total_rows
                complete[table] = True
                table_stats[table]["rows_exported"] = int(rows_written)
                table_stats[table]["status"] = "done"
                self._update_overall_progress(
                    done=done,
                    totals=totals,
                    complete=complete,
                    table_count=len(table_specs),
                    emit=_emit_progress,
                    message=f"{label}：完成",
                    stage="table_done",
                    counters={
                        "table": table,
                        "rows_total": total_rows,
                        "rows_exported": int(rows_written),
                        "format": table_stats[table]["format"],
                    },
                )

            duration_ms = int((time.perf_counter() - start_ts) * 1000)
            _emit_progress(
                100,
                "导出完成",
                stage="completed",
                counters={
                    "table_count": len(table_specs),
                    "file_count": len(created_files),
                    "csv_count": len(csv_exports),
                    "empty_tables": len(empty_tables),
                    "duration_ms": duration_ms,
                },
            )
            return {
                "case_id": case_id,
                "mode": mode_value,
                "export_format": format_value,
                "output_path": str(export_dir),
                "files": [str(path) for path in created_files],
                "csv_exports": [str(path) for path in csv_exports],
                "empty_tables": empty_tables,
                "table_stats": table_stats,
                "duration_ms": duration_ms,
                "naming": {
                    "folder_name": folder_name,
                    "base_name": safe_name,
                    "suffix": EXPORT_FOLDER_SUFFIX,
                    "tag": export_tag,
                },
            }
        except ExportCancelled:
            for path in created_files:
                try:
                    if path.exists():
                        path.unlink()
                except Exception:
                    pass
            try:
                if export_dir.exists() and not any(export_dir.iterdir()):
                    export_dir.rmdir()
            except Exception:
                pass
            raise
        finally:
            if con is not None:
                try:
                    con.close()
                except Exception:
                    pass

    @staticmethod
    def _resolve_requested_tables(
        filters: Optional[dict],
        label_to_table: Dict[str, str],
        table_to_label: Dict[str, str],
    ) -> List[str]:
        if not isinstance(filters, dict):
            return [table for table, _ in EXPORT_TABLES]
        raw_tables = filters.get("tables")
        if not isinstance(raw_tables, list) or not raw_tables:
            return [table for table, _ in EXPORT_TABLES]
        selected: List[str] = []
        for item in raw_tables:
            token = str(item or "").strip()
            if not token:
                continue
            if token in table_to_label:
                selected.append(token)
                continue
            if token in label_to_table:
                selected.append(label_to_table[token])
        seen: set[str] = set()
        ordered: List[str] = []
        for table, _ in EXPORT_TABLES:
            if table in selected and table not in seen:
                seen.add(table)
                ordered.append(table)
        return ordered

    @staticmethod
    def _table_exists(con: DuckDBEngine, table: str) -> bool:
        rows = con.query(
            "SELECT 1 FROM information_schema.tables WHERE table_schema='main' AND table_name=? LIMIT 1",
            (table,),
        )
        return bool(rows)

    @staticmethod
    def _resolve_table_name(con: DuckDBEngine, table: str) -> str:
        if table.startswith("fc_"):
            normalized = norm_table_name(table)
            if ExportRepository._table_exists(con, normalized):
                return normalized
        return table

    @staticmethod
    def _has_case_id(con: DuckDBEngine, table: str) -> bool:
        try:
            rows = con.query(
                "SELECT column_name FROM information_schema.columns WHERE table_schema='main' AND table_name=?",
                (table,),
            )
            return any(row[0] == "case_id" for row in rows)
        except Exception:
            return False

    @staticmethod
    def _count_rows(con: DuckDBEngine, table: str, has_case: bool, case_id: Optional[str] = None) -> int:
        if has_case and case_id:
            rows = con.query(f"SELECT COUNT(1) FROM {table} WHERE case_id=?", (case_id,))
        else:
            rows = con.query(f"SELECT COUNT(1) FROM {table}")
        row = rows[0] if rows else None
        return int(row[0] or 0) if row else 0

    @staticmethod
    def _table_columns(
        *,
        con: DuckDBEngine,
        table: str,
        base_table: str,
        use_cleaned: bool,
        case_id: str,
    ) -> tuple[list[str], list[str], list[str], dict[str, str]]:
        schema = FC_SCHEMAS.get(base_table)
        db_cols: list[str] = []
        type_map: dict[str, str] = {}
        try:
            rows = con.query(
                "SELECT column_name, data_type "
                "FROM information_schema.columns "
                "WHERE table_schema='main' AND table_name=? "
                "ORDER BY ordinal_position",
                (table,),
            )
            db_cols = [row[0] for row in rows]
            type_map = {row[0]: row[1] for row in rows}
        except Exception:
            db_cols = []
            type_map = {}

        cols: list[str] = []
        headers: list[str] = []
        if schema is not None:
            for header in schema.headers:
                col = schema.col_map.get(header)
                if col and (not db_cols or col in db_cols):
                    cols.append(col)
                    headers.append(header)
                    type_map.setdefault(col, "TEXT")
            if not headers and schema.headers:
                cols = [schema.col_map.get(header, header) for header in schema.headers]
                headers = list(schema.headers)
                for col in cols:
                    type_map.setdefault(col, "TEXT")

        if not cols and db_cols:
            cols = list(db_cols)
            headers = list(db_cols)

        if base_table == "fc_transaction" and use_cleaned:
            acct_info_map = [
                ("account_open_name", "账户开户名称"),
                ("opener_id_no", "开户人证件号码"),
            ]
            used_headers = set(headers)
            try:
                insert_at = headers.index("交易账号") + 1
            except ValueError:
                insert_at = len(headers)
            for col_name, header in acct_info_map:
                if header in used_headers:
                    continue
                cols.insert(insert_at, col_name if col_name in db_cols else f"'' AS {col_name}")
                headers.insert(insert_at, header)
                used_headers.add(header)
                type_map.setdefault(col_name, "TEXT")
                insert_at += 1

            extra_map = [
                ("clean_invalid", "无效数据"),
                ("clean_failed", "失败标记"),
                ("clean_reversal", "冲正标记"),
                ("clean_duplicate", "重复数据"),
            ]
            for col_name, header in extra_map:
                if col_name in db_cols:
                    cols.append(f"CASE WHEN COALESCE({col_name},0)=1 THEN '是' ELSE '' END AS {col_name}")
                else:
                    cols.append(f"'' AS {col_name}")
                headers.append(header)
                type_map.setdefault(col_name, "TEXT")

        if base_table in {"fc_transaction", "fc_account", "fc_sub_account"}:
            extra_keys = ExportRepository._extra_json_keys(
                con=con,
                table=table,
                db_cols=db_cols,
                case_id=case_id,
            )
            if extra_keys:
                used_headers = set(headers)
                for idx, key in enumerate(extra_keys, start=1):
                    alias = f"extra_{idx}"
                    path = ExportRepository._json_path(str(key))
                    cols.append(f"json_extract_string(extra_json, '{path}') AS {alias}")
                    headers.append(ExportRepository._unique_header(str(key), used_headers))
                    type_map.setdefault(alias, "TEXT")
                cols.append("extra_json")
                headers.append("扩展字段(原始)")
                type_map.setdefault("extra_json", "TEXT")
            elif "extra_json" in db_cols:
                cols.append("extra_json")
                headers.append("扩展字段")
                type_map.setdefault("extra_json", "TEXT")

        return cols, headers, db_cols, type_map

    @staticmethod
    def _extra_json_keys(
        *,
        con: DuckDBEngine,
        table: str,
        db_cols: list[str],
        case_id: str,
    ) -> list[str]:
        if "extra_json" not in db_cols:
            return []
        try:
            if ExportRepository._has_case_id(con, table):
                rows = con.query(
                    "SELECT DISTINCT k FROM ("
                    "  SELECT unnest(json_keys(extra_json)) AS k "
                    f"  FROM {table} "
                    "  WHERE case_id=? AND extra_json IS NOT NULL AND extra_json<>''"
                    ") t ORDER BY k",
                    (case_id,),
                )
            else:
                rows = con.query(
                    "SELECT DISTINCT k FROM ("
                    "  SELECT unnest(json_keys(extra_json)) AS k "
                    f"  FROM {table} "
                    "  WHERE extra_json IS NOT NULL AND extra_json<>''"
                    ") t ORDER BY k",
                )
            keys = [row[0] for row in rows if row and row[0]]
            return keys[:100]
        except Exception:
            return []

    @staticmethod
    def _json_path(key: str) -> str:
        safe = (key or "").replace("\\", "\\\\").replace('"', '\\"')
        return f'$.\"{safe}\"'

    @staticmethod
    def _unique_header(name: str, used: set[str]) -> str:
        if name not in used:
            used.add(name)
            return name
        idx = 2
        while f"{name}_{idx}" in used:
            idx += 1
        final_name = f"{name}_{idx}"
        used.add(final_name)
        return final_name

    @staticmethod
    def _strip_key_suffix_expr(expr: str) -> str:
        base = f"regexp_replace(COALESCE({expr}, ''), '\\\\s+', '', 'g')"
        return (
            "CASE "
            f"WHEN {base} IS NULL OR {base}='' THEN NULL "
            f"WHEN instr({base}, '-') > 0 OR instr({base}, '_') > 0 THEN "
            f"  CASE WHEN (instr({base}, '_') > 0 AND (instr({base}, '-') = 0 OR instr({base}, '_') < instr({base}, '-'))) THEN "
            f"    CASE WHEN instr({base}, '_') > 1 THEN substr({base}, 1, instr({base}, '_') - 1) ELSE NULL END "
            f"  ELSE "
            f"    CASE WHEN instr({base}, '-') > 1 THEN substr({base}, 1, instr({base}, '-') - 1) ELSE NULL END "
            f"  END "
            f"ELSE {base} END"
        )

    @staticmethod
    def _cleaned_select_cols(cols: list[str], db_cols: list[str]) -> list[str]:
        clean_map = {
            "amount": "clean_amount",
            "balance": "clean_balance",
            "dc_flag": "clean_dc_flag",
            "card_no": "clean_card_no",
            "acct_no": "clean_acct_no",
        }
        selected: list[str] = []
        for col in cols:
            clean_col = clean_map.get(col)
            if clean_col and clean_col in db_cols:
                if col in ("card_no", "acct_no"):
                    base_expr = f"COALESCE(NULLIF({clean_col},''), {col})"
                    selected.append(f"{ExportRepository._strip_key_suffix_expr(base_expr)} AS {col}")
                    continue
                if col == "amount" and "clean_amt_failed" in db_cols:
                    selected.append(
                        "CASE WHEN clean_amt_failed=1 "
                        "  AND TRY_CAST(COALESCE(NULLIF(clean_amount,''), NULL) AS DOUBLE) IS NULL "
                        "  AND TRY_CAST(COALESCE(NULLIF(amount,''), NULL) AS DOUBLE) IS NULL "
                        "THEN '转换失败' "
                        f"ELSE COALESCE(NULLIF({clean_col},''), {col}) END AS {col}"
                    )
                elif col == "balance" and "clean_bal_failed" in db_cols:
                    selected.append(
                        "CASE WHEN clean_bal_failed=1 "
                        "  AND TRY_CAST(COALESCE(NULLIF(clean_balance,''), NULL) AS DOUBLE) IS NULL "
                        "  AND TRY_CAST(COALESCE(NULLIF(balance,''), NULL) AS DOUBLE) IS NULL "
                        "THEN '转换失败' "
                        f"ELSE COALESCE(NULLIF({clean_col},''), {col}) END AS {col}"
                    )
                else:
                    selected.append(f"COALESCE(NULLIF({clean_col},''), {col}) AS {col}")
            else:
                selected.append(col)
        return selected

    @staticmethod
    def _build_sql(
        *,
        table: str,
        cols: list[str],
        has_case: bool,
        db_cols: list[str],
        type_map: dict[str, str],
        base_table: str,
        case_id: str,
        use_cleaned: bool,
    ) -> tuple[str, tuple]:
        if not cols and db_cols:
            cols = list(db_cols)
        if cols:
            if use_cleaned and base_table in ("fc_transaction", "fc_account"):
                cols_sql = ", ".join(
                    ExportRepository._apply_sanitize(
                        ExportRepository._cleaned_select_cols(cols, db_cols),
                        type_map,
                    )
                )
            else:
                cols_sql = ", ".join(ExportRepository._apply_sanitize(cols, type_map))
        else:
            cols_sql = "*"
        sql = f"SELECT {cols_sql} FROM {table}"
        params: tuple = ()
        if has_case:
            sql += " WHERE case_id=?"
            params = (case_id,)
        return sql, params

    @staticmethod
    def _is_text_type(dtype: Optional[str]) -> bool:
        if not dtype:
            return False
        lowered = dtype.lower()
        return "char" in lowered or "text" in lowered or "string" in lowered

    @staticmethod
    def _sanitize_key_expr(expr: str) -> str:
        cleaned = _sanitize_text_expr(expr)
        compact = f"regexp_replace(COALESCE({cleaned}, ''), '\\\\s+', '', 'g')"
        return f"NULLIF(trim({compact}), '')"

    @staticmethod
    def _apply_sanitize(cols: list[str], type_map: dict[str, str]) -> list[str]:
        key_fields = {
            "card_no",
            "acct_no",
            "counterparty_acct",
            "counterparty_card",
            "counterparty_acct_no",
            "counterparty_card_no",
        }
        out: list[str] = []
        for col in cols:
            expr, alias = split_alias(col)
            dtype = type_map.get(alias)
            if ExportRepository._is_text_type(dtype):
                if alias in key_fields:
                    out.append(f"{ExportRepository._sanitize_key_expr(expr)} AS {alias}")
                else:
                    out.append(f"{_sanitize_text_expr(expr)} AS {alias}")
            else:
                if expr != alias:
                    out.append(f"{expr} AS {alias}")
                else:
                    out.append(expr)
        return out

    def _export_csv(
        self,
        *,
        con: DuckDBEngine,
        case_id: str,
        table: str,
        has_case: bool,
        cols: list[str],
        headers: list[str],
        db_cols: list[str],
        type_map: dict[str, str],
        base_table: str,
        use_cleaned: bool,
        file_path: Path,
        on_chunk: Callable[[int], None],
    ) -> int:
        if not cols and db_cols:
            cols = list(db_cols)
        if use_cleaned and base_table in ("fc_transaction", "fc_account"):
            selected_cols = self._apply_sanitize(self._cleaned_select_cols(cols, db_cols), type_map)
        else:
            selected_cols = self._apply_sanitize(cols, type_map) if cols else []
        copy_sql, copy_params = build_csv_copy_sql(
            table=table,
            selected_cols=selected_cols,
            headers=headers,
            has_case=has_case,
            case_id=case_id,
            file_path=file_path,
        )
        with con.connection_operation() as connection:
            rows = connection.execute(copy_sql, copy_params).fetchall()
        written = int(rows[0][0] or 0) if rows else 0
        on_chunk(written)
        return written

    def _export_xlsx(
        self,
        *,
        con: DuckDBEngine,
        case_id: str,
        table: str,
        has_case: bool,
        cols: list[str],
        headers: list[str],
        db_cols: list[str],
        type_map: dict[str, str],
        base_table: str,
        use_cleaned: bool,
        file_path: Path,
        on_chunk: Callable[[int], None],
    ) -> int:
        try:
            from openpyxl import Workbook
        except Exception as exc:
            raise RuntimeError("missing dependency: openpyxl") from exc

        sql, params = self._build_sql(
            table=table,
            cols=cols,
            has_case=has_case,
            db_cols=db_cols,
            type_map=type_map,
            base_table=base_table,
            case_id=case_id,
            use_cleaned=use_cleaned,
        )
        with con.connection_operation() as connection:
            cur = connection.execute(sql, params)
            header = headers or ([desc[0] for desc in cur.description] if cur.description else [])
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet(title="Sheet1")
            if header:
                sheet.append(header)

            written = 0
            while True:
                rows = cur.fetchmany(CHUNK_SIZE)
                if not rows:
                    break
                for row in rows:
                    sheet.append(row)
                written += len(rows)
                on_chunk(written)
        workbook.save(file_path)
        return written

    @staticmethod
    def _update_overall_progress(
        *,
        done: Dict[str, int],
        totals: Dict[str, int],
        complete: Dict[str, bool],
        table_count: int,
        emit: Callable[[int, str, Optional[str], Optional[dict]], None],
        message: str,
        stage: Optional[str] = None,
        counters: Optional[dict] = None,
    ) -> None:
        total_rows = sum(int(value or 0) for value in totals.values())
        if total_rows > 0:
            done_rows = sum(int(value or 0) for value in done.values())
            pct = int(done_rows / total_rows * 100)
        else:
            completed = sum(1 for value in complete.values() if value)
            pct = int(completed / max(table_count, 1) * 100)
        emit(pct, message, stage=stage, counters=counters or {})

    @staticmethod
    def _on_chunk(
        *,
        written: int,
        table: str,
        table_label: str,
        done: Dict[str, int],
        totals: Dict[str, int],
        complete: Dict[str, bool],
        table_count: int,
        emit: Callable[[int, str, Optional[str], Optional[dict]], None],
        check_cancel: Callable[[], None],
    ) -> None:
        check_cancel()
        done[table] = int(written)
        ExportRepository._update_overall_progress(
            done=done,
            totals=totals,
            complete=complete,
            table_count=table_count,
            emit=emit,
            message=f"正在导出{table_label}",
            stage="table_progress",
            counters={
                "table": table,
                "rows_exported": int(written),
                "rows_total": int(totals.get(table) or 0),
            },
        )
